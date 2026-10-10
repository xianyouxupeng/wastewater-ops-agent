"""文档解析与切分：md/txt 直读，pdf/docx 提取文本，按 800 字递归切分"""
from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", "。", "；", "，", " ", ""],
)


def parse_file(path: Path) -> str:
    """按后缀解析为纯文本"""
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    raise ValueError(f"不支持的文件类型: {suffix}")


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(path: Path) -> str:
    import docx
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # 规程类文档大量信息在表格里
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def split_text(text: str) -> list[str]:
    """切分成 chunk 列表，去掉空白块"""
    return [c.strip() for c in _splitter.split_text(text) if c.strip()]
