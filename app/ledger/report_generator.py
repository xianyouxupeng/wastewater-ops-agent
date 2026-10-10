"""台账生成：Jinja2 模板渲染 markdown → 简易渲染为 Word（python-docx）。

模板独立成 .j2 文件，各厂可自行改字段/格式，不动代码。
"""
from io import BytesIO
from pathlib import Path

from docx import Document
from jinja2 import Environment

TEMPLATE_PATH = Path(__file__).parent / "report_template.md.j2"

# 表单默认值：界面录入缺项时补空，模板里统一显示"—"
DEFAULTS = {
    "plant_name": "本厂", "date": "", "shift": "白班", "operator": "",
    "inlet_cod": "", "outlet_cod": "", "inlet_nh3n": "", "outlet_nh3n": "",
    "outlet_tn": "", "outlet_tp": "", "outlet_ss": "",
    "inlet_ph": "", "outlet_ph": "",
    "flow": "", "temp": "", "do": "", "mlss": "", "sv30": "", "svi": "",
    "pac": "", "power": "", "summary": "", "abnormal": "",
}


def render_markdown(data: dict) -> str:
    """表单数据 → 台账 markdown（Jinja2 渲染，模板可独立修改）"""
    env = Environment(keep_trailing_newline=True)
    template = env.from_string(TEMPLATE_PATH.read_text(encoding="utf-8"))
    return template.render(**{**DEFAULTS, **data})


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell


def markdown_to_docx(md: str, doc: "Document | None" = None) -> Document:
    """最小 markdown 渲染：标题/表格/无序列表/段落，覆盖台账模板所需子集"""
    doc = doc or Document()
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [[c.strip() for c in row.strip().strip("|").split("|")]
                    for row in block]
            rows = [r for r in rows if not set("".join(r)) <= set("-: ")]  # 跳过分隔行
            if rows:
                _add_table(doc, rows)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=0)
            i += 1
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
            i += 1
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
            i += 1
        elif line.strip():
            doc.add_paragraph(line)
            i += 1
        else:
            i += 1
    return doc


def generate_report(data: dict) -> bytes:
    """台账主入口：表单数据 → docx 字节流（供界面下载按钮）"""
    doc = markdown_to_docx(render_markdown(data))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
